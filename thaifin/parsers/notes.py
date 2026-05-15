"""Parse ``NOTES.DOC[X]`` into a markdown blob.

v0 deliberately does no structural extraction. The ``NOTES`` document is a
long narrative + many tables (segment reporting, accounting policies,
related-party transactions, contingencies) and the high-value numbers we
already capture from the financial statements. We keep the full markdown
in case downstream consumers want to LLM over it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

import docx

from thaifin.parsers._convert import ensure_docx


@dataclass(frozen=True)
class NotesRow:
    """One row in ``notes_text.parquet`` (one per filing)."""

    filing_id: str
    symbol: str
    period: str
    raw_text_md: str


def _document_to_markdown(doc) -> str:
    """Same paragraph + table layout as the auditor report parser.

    Kept inline (not shared with auditor_report) because the auditor
    parser may evolve different heading/table conventions later.
    """
    out: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading 1") or style.startswith("title"):
            out.append(f"# {text}")
        elif style.startswith("heading 2"):
            out.append(f"## {text}")
        elif style.startswith("heading 3"):
            out.append(f"### {text}")
        else:
            out.append(text)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            out.append("| " + " | ".join(cells) + " |")
        out.append("")
    return "\n\n".join(out)


def parse_notes(
    data: bytes,
    *,
    filing_id: str,
    symbol: str,
    period: str,
) -> NotesRow:
    """Parse one ``NOTES.DOC[X]`` payload into a markdown row."""
    with TemporaryDirectory(prefix="thaifin-notes-") as tmp:
        safe_stem = re.sub(r"[^A-Za-z0-9_.-]", "_", filing_id) or "notes"
        docx_path = ensure_docx(data, Path(tmp), stem=safe_stem)
        doc = docx.Document(str(docx_path))
        return NotesRow(
            filing_id=filing_id,
            symbol=symbol,
            period=period,
            raw_text_md=_document_to_markdown(doc),
        )


def row_to_record(row: NotesRow) -> dict:
    return {
        "filing_id": row.filing_id,
        "symbol": row.symbol,
        "period": row.period,
        "raw_text_md": row.raw_text_md,
    }


__all__ = ["NotesRow", "parse_notes", "row_to_record"]
