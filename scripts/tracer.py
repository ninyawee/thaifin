"""Tracer bullet for slice #14 of PRD #11.

End-to-end vertical slice: download the PTT FY2025 annual filing zip from
Thai SEC IDISC, parse the Cash flow sheet to extract the canonical CapEx
line, write a 1-row ``financial_lines.parquet`` in the tagged long-format
schema, and upload it to HuggingFace at revision ``tracer.0``.

Run via ``fnox exec`` so HF_TOKEN is injected only into the subprocess
environment and never lands in argv or shell history::

    fnox exec -- uv run python scripts/tracer.py

Idempotent: re-running uploads the same parquet content; HF deduplicates by
content hash but the commit + tag move forward.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import httpx
import openpyxl
import pandas as pd
from huggingface_hub import HfApi

# --- Constants from issue #14 -----------------------------------------------

PTT_2025_ZIP_URL = (
    "https://market.sec.or.th/public/idisc/Download"
    "?FILEID=dat/news/202602/0646FIN190220261747440265T.zip"
)
CAPEX_LABEL_TH = (
    "เงินสดจ่ายสำหรับที่ดิน อาคารและอุปกรณ์ และอสังหาริมทรัพย์เพื่อการลงทุน"
)
EXPECTED_CAPEX_VALUE = -159_512_958_954

HF_REPO_ID = "ninyawee/thaifin-financials"
HF_REVISION_TAG = "tracer.0"
HF_COMMIT_MESSAGE = "Tracer bullet: PTT FY2025 capex"

# --- Schema columns (must match docs/adr/0001-tagged-long-format.md) --------

SCHEMA_COLUMNS = [
    "symbol",
    "period",
    "statement",
    "concept",
    "raw_label_th",
    "value",
    "audit_basis",
    "consolidation",
    "filing_id",
]


def fetch_zip(url: str, dest: Path) -> Path:
    """Download a zip to ``dest`` and return its path."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    with httpx.Client(follow_redirects=True, timeout=60.0) as client:
        resp = client.get(url)
        resp.raise_for_status()
        dest.write_bytes(resp.content)
    return dest


def extract_zip(zip_path: Path, out_dir: Path) -> Path:
    """Unzip ``zip_path`` into ``out_dir`` and return the path to the XLS."""
    out_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as zf:
        zf.extractall(out_dir)
    candidates = list(out_dir.glob("FINANCIAL_STATEMENTS.*"))
    if not candidates:
        raise FileNotFoundError(f"No FINANCIAL_STATEMENTS.* in {out_dir}")
    return candidates[0]


def open_xls_as_xlsx(xls_path: Path) -> openpyxl.Workbook:
    """Open the SEC FINANCIAL_STATEMENTS file via openpyxl.

    SEC names the file ``.XLS`` regardless of whether it's legacy BIFF or
    OOXML. PTT's Feb-2026 filing is OOXML behind a ``.XLS`` extension.
    Openpyxl rejects by extension, so copy to a ``.xlsx`` sibling first.
    If the file turns out to be true legacy BIFF, fall back to libreoffice
    headless conversion (still produces .xlsx).
    """
    head = xls_path.read_bytes()[:4]
    if head[:2] == b"PK":
        # OOXML (zip-based); openpyxl can read it, just needs the right ext.
        target = xls_path.with_suffix(".xlsx")
        if target != xls_path:
            shutil.copyfile(xls_path, target)
        return openpyxl.load_workbook(target, data_only=True)
    # Legacy BIFF: convert via libreoffice. Argv-form subprocess (no shell).
    out_dir = xls_path.parent
    proc = subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            str(xls_path),
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        raise RuntimeError(
            f"libreoffice conversion failed for {xls_path} "
            f"(exit {proc.returncode}): {proc.stderr}"
        )
    target = out_dir / (xls_path.stem + ".xlsx")
    return openpyxl.load_workbook(target, data_only=True)


def find_capex_row(wb: openpyxl.Workbook, label: str) -> tuple[int, float]:
    """Locate the CapEx row by exact Thai label match in the Cash flow sheet.

    Returns ``(row_index, fy2025_consolidated_value)``. Raises if not found
    or if the value is missing.

    The PTT Cash flow sheet layout (Feb-2026 filing):
      - Section headers in column 2.
      - Line-item labels in column 3.
      - Periods in row 8: col 9 = consolidated current FY (2568=2025),
        col 11 = consolidated prior FY (2567=2024), col 13 = company
        current, col 15 = company prior.
    """
    if "Cash flow" not in wb.sheetnames:
        raise ValueError(f"Cash flow sheet missing; have: {wb.sheetnames}")
    ws = wb["Cash flow"]
    for row in range(1, ws.max_row + 1):
        for col in (2, 3):
            cell = ws.cell(row=row, column=col).value
            if isinstance(cell, str) and cell.strip() == label:
                value = ws.cell(row=row, column=9).value
                if value is None:
                    raise ValueError(
                        f"Found label at row {row} col {col} but consolidated "
                        f"FY value (col 9) is empty"
                    )
                return row, float(value)
    raise ValueError(f"Label not found in Cash flow: {label!r}")


def build_parquet(value: float, label: str, out_path: Path) -> Path:
    """Write a 1-row financial_lines parquet with the tracer schema."""
    row = {
        "symbol": "PTT",
        "period": "2025",
        "statement": "CF",
        "concept": "capex",
        "raw_label_th": label,
        "value": float(value),
        "audit_basis": "audited",
        "consolidation": "consolidated",
        "filing_id": "PTT-2025-A-tracer",
    }
    df = pd.DataFrame([row], columns=SCHEMA_COLUMNS)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(out_path, engine="pyarrow", index=False)
    return out_path


def publish_to_hf(parquet_path: Path) -> None:
    """Upload parquet to HF and tag the commit as ``tracer.0``."""
    token = os.environ.get("HF_TOKEN")
    if not token:
        raise RuntimeError(
            "HF_TOKEN not set; run via "
            "`fnox exec -- uv run python scripts/tracer.py`"
        )
    api = HfApi(token=token)
    api.upload_file(
        path_or_fileobj=str(parquet_path),
        path_in_repo="financial_lines.parquet",
        repo_id=HF_REPO_ID,
        repo_type="dataset",
        commit_message=HF_COMMIT_MESSAGE,
    )
    # Recreate tag if it already exists so the slice can be re-run.
    try:
        api.delete_tag(
            repo_id=HF_REPO_ID, tag=HF_REVISION_TAG, repo_type="dataset"
        )
    except Exception:
        pass
    api.create_tag(
        repo_id=HF_REPO_ID,
        tag=HF_REVISION_TAG,
        repo_type="dataset",
        tag_message=HF_COMMIT_MESSAGE,
    )


def main() -> None:
    work = Path("/tmp/thaifin-tracer")
    work.mkdir(parents=True, exist_ok=True)

    print(f"[1/4] Fetch PTT 2025 zip -> {work / 'ptt.zip'}", flush=True)
    zip_path = fetch_zip(PTT_2025_ZIP_URL, work / "ptt.zip")

    print(f"[2/4] Extract zip -> {work / 'extracted'}", flush=True)
    xls_path = extract_zip(zip_path, work / "extracted")
    wb = open_xls_as_xlsx(xls_path)

    print("[3/4] Find CapEx row by label match", flush=True)
    row_idx, value = find_capex_row(wb, CAPEX_LABEL_TH)
    print(f"      row {row_idx}: value = {value:,.0f} THB", flush=True)
    if int(value) != EXPECTED_CAPEX_VALUE:
        raise AssertionError(
            f"Parsed value {int(value)} != expected {EXPECTED_CAPEX_VALUE}"
        )

    parquet_path = work / "financial_lines.parquet"
    build_parquet(value, CAPEX_LABEL_TH, parquet_path)
    print(
        f"      wrote {parquet_path} ({parquet_path.stat().st_size} bytes)",
        flush=True,
    )

    if "--no-upload" in sys.argv:
        print("[4/4] --no-upload set; skipping HF push", flush=True)
        return

    print(f"[4/4] Upload to HF {HF_REPO_ID} @ {HF_REVISION_TAG}", flush=True)
    publish_to_hf(parquet_path)
    print("      done.", flush=True)


if __name__ == "__main__":
    main()
